import re
from typing import List
from elements import ListElement, Paragraph, UndefinedElement
from elements import Element, Formula, Header, Image, Listing, Table
import re
from docx import Document
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.oxml.ns import qn, nsmap
from lxml import etree

from tree_numbering import TreeNumbering

# Вспомогательные функции
def is_header(paragraph):
    return paragraph.style.name.startswith('Heading')

def get_header_level(paragraph):
    return int(paragraph.style.name.split()[1]) if paragraph.style.name.startswith('Heading') else 1

def is_list_item(paragraph):
    if paragraph.style is None:
        return False
    style_name = paragraph.style.name.lower()
    return 'list' in style_name or 'bullet' in style_name or 'number' in style_name

def get_list_type(paragraph):
    if paragraph.style is None:
        return 'unordered'
    style_name = paragraph.style.name.lower()
    if 'bullet' in style_name:
        return 'unordered'
    elif 'number' in style_name:
        return 'ordered'
    return 'unordered'

def is_table_caption(paragraph):
    text = paragraph.text.strip()
    pattern = r'^Таблица \d+\.\d+ – .+$'
    return bool(re.fullmatch(pattern, text))

def is_image_caption(paragraph):
    text = paragraph.text.strip()
    pattern = r'^Рисунок \d+\.\d+ – .+$'
    return bool(re.fullmatch(pattern, text))

def is_listing_caption(paragraph):
    text = paragraph.text.strip()
    pattern = r'^Листинг \d+\.\d+ – .+$'
    return bool(re.fullmatch(pattern, text))

def is_formula_caption(paragraph):
    text = paragraph.text.strip()
    pattern = r'^\(\d+\.\d+\)$'
    return bool(re.fullmatch(pattern, text))

def has_image(paragraph):
    has_pic = any(run.element.tag.endswith('pic') for run in paragraph.runs)
    has_inline = paragraph._element.findall('.//wp:inline', namespaces=nsmap)
    # print(f"Checking paragraph for image: has_pic={has_pic}, has_inline={len(has_inline) > 0}")
    return has_pic or has_inline

def extract_image(paragraph):
    for run in paragraph.runs:
        if run.element.tag.endswith('pic'):
            blip = run.element.find('.//a:blip', namespaces=nsmap)
            if blip is not None:
                rId = blip.get(qn('r:embed'))
                if rId and rId in paragraph.part.related_parts:
                    image_part = paragraph.part.related_parts[rId]
                    return image_part.blob
    for inline in paragraph._element.findall('.//wp:inline', namespaces=nsmap):
        blip = inline.find('.//a:blip', namespaces=nsmap)
        if blip is not None:
            rId = blip.get(qn('r:embed'))
            if rId and rId in paragraph.part.related_parts:
                image_part = paragraph.part.related_parts[rId]
                return image_part.blob
    return b''

def has_formula(paragraph):
    has_omath = paragraph._element.findall('.//m:oMath', namespaces=nsmap)
    has_omathpara = paragraph._element.findall('.//m:oMathPara', namespaces=nsmap)
    # print(f"Checking paragraph for formula: has_omath={len(has_omath) > 0}, has_omathpara={len(has_omathpara) > 0}")
    return has_omath or has_omathpara

def extract_formula_and_caption(paragraph):
    formula_xml = []
    omath = paragraph._element.findall('.//m:oMath', namespaces=nsmap)
    omathpara = paragraph._element.findall('.//m:oMathPara', namespaces=nsmap)
    for elem in omath + omathpara:
        formula_xml.append(etree.tostring(elem, encoding='unicode'))
    formula = ''.join(formula_xml)
    caption = None
    text = paragraph.text.strip()
    caption_match = re.search(r'\(\d+\.\d+\)', text)
    if caption_match:
        caption = caption_match.group(0)
    # print(f"Extracting formula: formula_present={bool(formula)}, caption={caption}")
    return formula, caption

def is_listing(paragraph):
    # Проверяем наличие границ (w:pBdr) в свойствах параграфа
    pPr = paragraph._element.find('.//w:pPr', namespaces=nsmap)
    has_borders = pPr is not None and pPr.find('.//w:pBdr', namespaces=nsmap) is not None
    # Проверяем стиль Listing как запасной вариант
    is_listing_style = paragraph.style.name == 'Listing'
    # print(f"Checking listing: has_borders={has_borders}, is_listing_style={is_listing_style}")
    return has_borders or is_listing_style

# Основная функция парсинга
def parse_document(doc_path):
    doc = Document(doc_path)
    tree = []
    stack = [tree]
    current_level = 0
    pending_table_caption = None
    paragraph_index = 0
    table_index = 0
    elements = list(doc.element.body)

    i = 0
    while i < len(elements):
        element = elements[i]
        if element.tag.endswith('p'):
            if paragraph_index >= len(doc.paragraphs):
                parent = stack[-1]
                parent.append(UndefinedElement(str(element), element))
                i += 1
                continue
            paragraph = doc.paragraphs[paragraph_index]
            paragraph_index += 1
            if is_header(paragraph):
                level = get_header_level(paragraph)
                text = paragraph.text.strip()
                if not text:
                    i += 1
                    continue
                while len(stack) > 1 and current_level >= level:
                    stack.pop()
                    current_level -= 1
                parent = stack[-1]
                new_header = Header(level, text, [], paragraph, "")
                parent.append(new_header)
                stack.append(new_header.children)
                current_level = level
                i += 1
            elif is_list_item(paragraph):
                list_type = get_list_type(paragraph)
                list_items = [paragraph.text.strip() or "(empty)"]
                list_paragraphs = [paragraph]
                j = i + 1
                next_paragraph_index = paragraph_index
                while (j < len(elements) and elements[j].tag.endswith('p') and 
                       next_paragraph_index < len(doc.paragraphs) and
                       is_list_item(doc.paragraphs[next_paragraph_index]) and 
                       get_list_type(doc.paragraphs[next_paragraph_index]) == list_type):
                    next_text = doc.paragraphs[next_paragraph_index].text.strip()
                    list_items.append(next_text or "(empty)")
                    list_paragraphs.append(doc.paragraphs[next_paragraph_index])
                    next_paragraph_index += 1
                    j += 1
                parent = stack[-1]
                parent.append(ListElement(list_items, list_type, list_paragraphs))
                i = j
                paragraph_index = next_paragraph_index
            elif is_table_caption(paragraph):
                pending_table_caption = (paragraph.text.strip(), paragraph)
                i += 1
            elif has_image(paragraph):
                image_data = extract_image(paragraph)
                caption = None
                caption_paragraph = None
                if (i + 1 < len(elements) and elements[i + 1].tag.endswith('p') and 
                    paragraph_index < len(doc.paragraphs) and
                    is_image_caption(doc.paragraphs[paragraph_index])):
                    caption = doc.paragraphs[paragraph_index].text.strip()
                    caption_paragraph = doc.paragraphs[paragraph_index]
                    paragraph_index += 1
                    i += 1
                parent = stack[-1]
                parent.append(Image(image_data, caption, paragraph, caption_paragraph, ""))
                i += 1
            elif has_formula(paragraph):
                formula, caption = extract_formula_and_caption(paragraph)
                caption_paragraph = None
                if not caption:  # Check next paragraph only if no caption in current
                    if (i + 1 < len(elements) and elements[i + 1].tag.endswith('p') and 
                        paragraph_index < len(doc.paragraphs) and
                        is_formula_caption(doc.paragraphs[paragraph_index])):
                        caption = doc.paragraphs[paragraph_index].text.strip()
                        caption_paragraph = doc.paragraphs[paragraph_index]
                        paragraph_index += 1
                        i += 1
                if formula:
                    parent = stack[-1]
                    parent.append(Formula(formula, caption, paragraph, caption_paragraph, ""))
                else:
                    parent = stack[-1]
                    parent.append(UndefinedElement(str(paragraph._element), paragraph))
                i += 1
            elif is_listing(paragraph):
                listing_texts = [paragraph.text.strip() or "(empty)"]
                listing_paragraphs = [paragraph]
                j = i + 1
                next_paragraph_index = paragraph_index
                while (j < len(elements) and elements[j].tag.endswith('p') and 
                       next_paragraph_index < len(doc.paragraphs) and
                       is_listing(doc.paragraphs[next_paragraph_index]) and
                       not is_listing_caption(doc.paragraphs[next_paragraph_index])):
                    next_text = doc.paragraphs[next_paragraph_index].text.strip()
                    listing_texts.append(next_text or "(empty)")
                    listing_paragraphs.append(doc.paragraphs[next_paragraph_index])
                    next_paragraph_index += 1
                    j += 1
                text = '\n'.join(listing_texts)
                caption = None
                caption_paragraph = None
                if (j < len(elements) and elements[j].tag.endswith('p') and 
                    next_paragraph_index < len(doc.paragraphs) and
                    is_listing_caption(doc.paragraphs[next_paragraph_index])):
                    caption = doc.paragraphs[next_paragraph_index].text.strip()
                    caption_paragraph = doc.paragraphs[next_paragraph_index]
                    next_paragraph_index += 1
                    j += 1
                parent = stack[-1]
                parent.append(Listing(text, caption, listing_paragraphs, caption_paragraph, ""))
                i = j
                paragraph_index = next_paragraph_index
            else:
                text = paragraph.text.strip()
                if text:
                    parent = stack[-1]
                    parent.append(Paragraph(text, paragraph))
                i += 1
        elif element.tag.endswith('tbl'):
            if table_index < len(doc.tables):
                table = doc.tables[table_index]
                table_index += 1
                caption, caption_paragraph = pending_table_caption if pending_table_caption else (None, None)
                pending_table_caption = None
                data = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                parent = stack[-1]
                parent.append(Table(data, caption, table, caption_paragraph, ""))
            else:
                parent = stack[-1]
                parent.append(UndefinedElement(str(element), element))
            i += 1
        else:
            parent = stack[-1]
            parent.append(UndefinedElement(str(element), element))
            i += 1

    return tree


# Функция для красивого вывода дерева
def print_tree(elements, indent=0, prefix=""):
    for i, elem in enumerate(elements):
        is_last = i == len(elements) - 1
        branch = "└── " if is_last else "├── "
        indent_str = prefix + branch
        content = ""
        
        if isinstance(elem, Header):
            content = f"Header (Level {elem.level}, {elem.text_number_of_element}): {elem.text[:50]}{'...' if len(elem.text) > 50 else ''}"
        elif isinstance(elem, Paragraph):
            content = f"Paragraph: {elem.text[:50]}{'...' if len(elem.text) > 50 else ''}"
        elif isinstance(elem, ListElement):
            first_item = elem.items[0][:47] + '...' if len(elem.items) > 0 and len(elem.items[0]) > 47 else elem.items[0] if len(elem.items) > 0 else 'No items'
            content = f"List ({elem.list_type}, {len(elem.items)} items): {first_item}"
        elif isinstance(elem, Image):
            image_size = f"{len(elem.image_data)} bytes" if elem.image_data else "No data"
            content = f"Image ({elem.text_number_of_element}): {elem.caption[:44] if elem.caption else 'No caption'} ({image_size})"
            content = content[:50] + '...' if len(content) > 50 else content
        elif isinstance(elem, Table):
            first_cell = elem.data[0][0][:41] if elem.data and elem.data[0] else 'No data'
            content = f"Table ({elem.text_number_of_element}, {len(elem.data)} rows): {elem.caption[:35] if elem.caption else 'No caption'}, {first_cell}"
            content = content[:50] + '...' if len(content) > 50 else content
        elif isinstance(elem, Formula):
            formula_preview = elem.formula[:41] if elem.formula else 'No formula'
            content = f"Formula ({elem.text_number_of_element}): {elem.caption[:35] if elem.caption else 'No caption'}, {formula_preview}"
            content = content[:50] + '...' if len(content) > 50 else content
        elif isinstance(elem, Listing):
            listing_preview = elem.text[:41] if elem.text else 'No text'
            content = f"Listing ({elem.text_number_of_element}): {elem.caption[:35] if elem.caption else 'No caption'}, {listing_preview}"
            content = content[:50] + '...' if len(content) > 50 else content
        elif isinstance(elem, UndefinedElement):
            content = f"Undefined: {elem.content[:50]}{'...' if len(elem.content) > 50 else ''}"
        
        print(f"{indent_str}{content}")
        
        if isinstance(elem, Header) and elem.children:
            new_prefix = prefix + ("    " if is_last else "│   ")
            print_tree(elem.children, indent + 1, new_prefix)

# Пример использования
if __name__ == "__main__":
    doc_path = "example.docx"
    elements_tree = parse_document(doc_path)
    numberer = TreeNumbering(elements_tree)
    numberer.assign_numbers()
    print("Document Structure:")
    print_tree(elements_tree)
    existing_image = elements_tree[1].children[3].children[2]  # "Рисунок 3.2"
    new_image = Image(b'image_data', "New image caption", None, None, "")
    pa = Paragraph(b'image_data test 32', None)
    numberer.add_after(existing_image, pa)
    numberer.add_after(pa, new_image)
    new_header = Header(level=1, text="1", children=[])
    new_header1 = Header(level=1, text="2", children=[])
    new_header2 = Header(level=1, text="3", children=[])
    numberer.add_top_level_header(new_header)  # App
    numberer.add_top_level_header(new_header1)  # App
    after_header = elements_tree[0]  # An existing top-level Header
    numberer.add_top_level_header(new_header2, after_header=new_header)
    numberer = TreeNumbering(elements_tree)

   
    pda = Paragraph(b'image_data test 32', None)
    new_imagdse = Image(b'image_data', "New image caption", None, None, "")

    numberer.add_after(new_header2, pda)
    numberer.add_after(new_header2, new_imagdse)
    new_headeriner = Header(level=2, text="dsa1", children=[])
    new_headeriner_one = Header(level=3, text="dsa1", children=[])

    numberer.add_nested_header(new_header2,new_headeriner)
    numberer.add_nested_header(new_headeriner,new_headeriner_one)
    new_imagdseds = Image(b'image_data', "New image caption", None, None, "")
    numberer.add_after(new_headeriner,new_imagdseds)
    new_headerdasd = Header(level=1, text="new", children=[])
    numberer.add_top_level_header(new_headerdasd,new_header)  # App

    print_tree(elements_tree)