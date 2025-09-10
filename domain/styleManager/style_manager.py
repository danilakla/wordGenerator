
from typing import List
from docx import Document
from elements.element import Element
from elements.header import Header
from elements.formula import Formula
from elements.image import Image
from elements.list_element import ListElement
from elements.listing import Listing
from elements.table import Table
from elements.paragraph import Paragraph
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_PARAGRAPH_ALIGNMENT,WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.oxml import parse_xml

from util.config_service import ConfigService
from util.neighbor_finder import NeighborFinder
class StyleManager:
    
    def __init__(self, doc: Document,tree: List[Element],univer:str):
        self.tree = tree
        self.COLOR_MAP = {
    "black": RGBColor(0, 0, 0),
    "white": RGBColor(255, 255, 255),
    "red": RGBColor(255, 0, 0),
    "green": RGBColor(0, 255, 0),
    "blue": RGBColor(0, 0, 255),
    # Add more as needed
}
        
        self.WD_LINE_SPACING_MAP = {
    "single":WD_LINE_SPACING.SINGLE,
    "Multiple": WD_LINE_SPACING.MULTIPLE,
    "Double": WD_LINE_SPACING.DOUBLE,
    "AtLeast": WD_LINE_SPACING.AT_LEAST,
    "Exactyl": WD_LINE_SPACING.EXACTLY
    # Add more as needed
}
        
        self.WD_ALIGN_PARAGRAPH_MAP = {
    "center":WD_ALIGN_PARAGRAPH.CENTER,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY
    # Add more as needed
}
        
        self.WD_PARAGRAPH_ALIGNMENT_MAP = {
    "center":WD_PARAGRAPH_ALIGNMENT.CENTER,
    "left": WD_PARAGRAPH_ALIGNMENT.LEFT,
    "right": WD_PARAGRAPH_ALIGNMENT.RIGHT,
    "justify": WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    # Add more as needed
}
        self.WD_CELL_VERTICAL_ALIGNMENT_MAP = {
    "center":WD_CELL_VERTICAL_ALIGNMENT.CENTER,
    "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
    "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
    "both": WD_CELL_VERTICAL_ALIGNMENT.BOTH    # Add more as needed
}
        self.config_service = ConfigService("resource/"+univer+"formatting_rules.json")
        self.neighbor_finder=NeighborFinder(self.tree)
        self.doc = doc
        self._ensure_styles()
        

    def _ensure_styles(self):
        """Ensures required styles exist in the document."""
        settings = self.doc.settings
        math_pr = settings.element.find(qn('w:mathPr'))
        if math_pr is None:
            math_pr = OxmlElement('w:mathPr')
            settings.element.append(math_pr)

        math_font = math_pr.find(qn('m:mathFont'))
        if math_font is None:
            math_font = OxmlElement('m:mathFont')
            math_pr.append(math_font)

        math_font.set(qn('m:val'), "XITS Math")
        
        for section in self.doc.sections:
            section.top_margin = Cm(self.config_service.get("documentFormatting.general.margins.top"))
            section.bottom_margin = Cm(self.config_service.get("documentFormatting.general.margins.bottom"))
            section.left_margin = Cm(self.config_service.get("documentFormatting.general.margins.left"))
            section.right_margin = Cm(self.config_service.get("documentFormatting.general.margins.right"))
    
        auto_hyphenation = settings.element.find(qn('w:autoHyphenation'))
        if auto_hyphenation is None:
            auto_hyphenation = parse_xml(r'<w:autoHyphenation xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="true"/>')
            settings.element.append(auto_hyphenation)
        else:
            auto_hyphenation.set(qn('w:val'), 'true')

        # Устанавливаем параметры переносов
        hyphenation_zone = settings.element.find(qn('w:hyphenationZone'))
        if hyphenation_zone is None:
            hyphenation_zone = parse_xml(r'<w:hyphenationZone xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="360"/>')
            settings.element.append(hyphenation_zone)

        consecutive_hyphens = settings.element.find(qn('w:consecutiveHyphenLimit'))
        if consecutive_hyphens is None:
            consecutive_hyphens = parse_xml(r'<w:consecutiveHyphenLimit xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="3"/>')
            settings.element.append(consecutive_hyphens)
   
    def cm_to_pt(self,cm):
        """Конвертирует сантиметры в пункты (1 см ≈ 28.35 pt)."""
        return Pt(cm * 28.3464567)
    def apply_header_style(self, header:Header):
        if  header.docx_element:
            for run in header.docx_element.runs:
                run.font.name = self.config_service.get("documentFormatting.headings.font.family")  
                run.font.size = Pt(self.config_service.get("documentFormatting.headings.font.size"))   
                run.font.bold = self.config_service.get("documentFormatting.headings.font.style.bold")       
                run.font.italic = self.config_service.get("documentFormatting.headings.font.style.italic")  
                run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.headings.font.color")] 
                if header.level==1 and self.config_service.get("documentFormatting.headings.font.uppercase",False):
                    run.text= run.text.upper()

            paragraph_format  = header.docx_element.paragraph_format
            paragraph_format.keep_together=False
            paragraph_format.keep_with_next =False
            paragraph_format.widow_control = False
            paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.headings.font.lineSpacing", 0)]
            if paragraph_format.line_spacing_rule is WD_LINE_SPACING.EXACTLY:
                paragraph_format.line_spacing=Pt(self.config_service.get("documentFormatting.general.font.line_space_count", 18))
            paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.headings.section.paragraph_format.indent_left", 0))      
            paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.headings.section.paragraph_format.indent_right", 0))   
            paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.headings.section.paragraph_format.first_line", 0))
            neighbor_before = self.neighbor_finder.get_before(header)
            if header.level==1:
                paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.headings.section.paragraph_format.spacingBefore", 0))  # Space before paragraph
                paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.headings.section.paragraph_format.spacingAfter", 0))
                paragraph_format.alignment =self.WD_ALIGN_PARAGRAPH_MAP[self.config_service.get("documentFormatting.headings.section.paragraph_format.alignment", 0)]
              
            elif header.level>1:
                if isinstance(neighbor_before,Header):
                    neighbor_before.docx_element.paragraph_format.space_after=Pt(self.config_service.get("documentFormatting.headings.subsection.spacingBeforeAnotherSubSection", 0))  # Space before paragraph
                    paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.headings.subsection.spacingBeforeAnotherSubSection", 0))  # Space before paragraph
                    paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.headings.subsection.spacingAfter", 0))
                else:
                    paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.headings.subsection.spacingBefore", 0))  # Space before paragraph
                    paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.headings.subsection.spacingAfter", 0))

    def apply_image_style(self, image: Image):
        for run in image.docx_caption_paragraph.runs:
            run.font.name = self.config_service.get("documentFormatting.figures.family")  
            run.font.size = Pt(self.config_service.get("documentFormatting.figures.size"))   
            run.font.bold = self.config_service.get("documentFormatting.figures.style.bold")       
            run.font.italic = self.config_service.get("documentFormatting.figures.style.italic")  
            run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.figures.color")] 
        
        paragraph_format  = image.docx_caption_paragraph.paragraph_format
        paragraph_format.keep_together=False
        paragraph_format.keep_with_next =False
        paragraph_format.widow_control = False
        if paragraph_format.line_spacing_rule is WD_LINE_SPACING.EXACTLY:
            paragraph_format.line_spacing=Pt(self.config_service.get("documentFormatting.general.font.line_space_count", 18))
        paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.figures.lineSpacing", 0)]
        paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.figures.paragraph_format.indent_left", 0))      
        paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.figures.paragraph_format.indent_right", 0))   
        paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.figures.paragraph_format.first_line", 0))
        paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.figures.paragraph_format.spacingBefore", 0))  # Space before paragraph
        paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.figures.paragraph_format.spacingAfter", 0))
        paragraph_format.alignment=self.WD_ALIGN_PARAGRAPH_MAP[self.config_service.get("documentFormatting.figures.caption.alignment", 0)]

        paragraph_format  = image.docx_image_paragraph.paragraph_format
        paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.figures.lineSpacing", 0)]
        paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.figures.paragraph_format.indent_left", 0))      
        paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.figures.paragraph_format.indent_right", 0))   
        paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.figures.paragraph_format.first_line", 0))
        paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.figures.spacingBefore", 0))  # Space before paragraph
        paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.figures.spacingAfter", 0))
        paragraph_format.alignment=self.WD_ALIGN_PARAGRAPH_MAP[self.config_service.get("documentFormatting.figures.alignment", 0)]

    def appply_move_table(self,table):
        neighbor_after = self.neighbor_finder.get_after(table)
        if isinstance(neighbor_after,Paragraph):
            neighbor_after.docx_element.paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
        if isinstance(neighbor_after,ListElement):
            neighbor_after.docx_elements[0].paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
    def apply_table_after_update_run(self,table:Table):
        for row, text_row in zip(table.docx_table.rows, table.data):
            for cell, cell_object_tab in zip(row.cells, text_row):
                for run in cell.paragraphs[0].runs:
                    run.font.name = self.config_service.get("documentFormatting.tables.family")  
                    run.font.size = Pt(self.config_service.get("documentFormatting.tables.size"))   
                    run.font.bold = self.config_service.get("documentFormatting.tables.bold")       
                    run.font.italic = self.config_service.get("documentFormatting.tables.style.italic")  
                    run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.tables.color")  ] 
            
    def apply_table_style(self, table: Table):
        section = self.doc.sections[0]
        page_width = section.page_width - section.left_margin - section.right_margin
        tb_l=table.docx_table
        tb_l.width = page_width

        tb_l.autofit = False  # Фиксированные ширины
        tb_l.allow_autofit = False
        neighbor_after = self.neighbor_finder.get_after(table)
        if isinstance(neighbor_after,Paragraph):
            neighbor_after.docx_element.paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
        if isinstance(neighbor_after,ListElement):
            neighbor_after.docx_elements[0].paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
        for run in table.docx_caption_paragraph.runs:
            run.font.name = self.config_service.get("documentFormatting.tables.caption.family")  
            run.font.size = Pt(self.config_service.get("documentFormatting.tables.caption.size"))   
            run.font.bold = self.config_service.get("documentFormatting.tables.caption.bold")       
            run.font.italic = self.config_service.get("documentFormatting.tables.caption.style.italic")  
            run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.tables.caption.color")] 
        paragraph_format  = table.docx_caption_paragraph.paragraph_format

        paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.tables.caption.paragraph_format.line_sp", 0)]
        paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.tables.caption.paragraph_format.indent_left", 0))      
        paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.tables.caption.paragraph_format.indent_right", 0))   
        paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.tables.caption.paragraph_format.first_line", 0))
        paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.tables.caption.paragraph_format.spacingBefore", 0))  # Space before paragraph
        paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.tables.caption.paragraph_format.spacingAfter", 0))
        paragraph_format.keep_together=False
        paragraph_format.keep_with_next =False
        paragraph_format.widow_control = False

        paragraph_format.alignment=self.WD_ALIGN_PARAGRAPH_MAP[self.config_service.get("documentFormatting.tables.caption.paragraph_format.alignment", 0)]


        for row, text_row in zip(table.docx_table.rows, table.data):
            for cell, cell_object_tab in zip(row.cells, text_row):
                for run in cell.paragraphs[0].runs:
                    run.font.name = self.config_service.get("documentFormatting.tables.family")  
                    run.font.size = Pt(self.config_service.get("documentFormatting.tables.size"))   
                    run.font.bold = self.config_service.get("documentFormatting.tables.bold")       
                    run.font.italic = self.config_service.get("documentFormatting.tables.style.italic")  
                    run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.tables.color")  ] 
                paragraph_format  = cell.paragraphs[0].paragraph_format

                paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.tables.paragraph_format.line_sp", 0)]
                paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.indent_left", 0))      
                paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.indent_right", 0))   
                paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.tables.paragraph_format.first_line", 0))
                paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.spacingBefore", 0))  # Space before paragraph
                paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.spacingAfter", 0))
                paragraph_format.alignment=cell_object_tab.horizontal
                cell.vertical_alignment= cell_object_tab.vertical
                border_style = {
                    'top': {'val': 'single', 'sz': '4', 'color': '000000'},
                    'left': {'val': 'single', 'sz': '4', 'color': '000000'},
                    'bottom': {'val': 'single', 'sz': '4', 'color': '000000'},
                    'right': {'val': 'single', 'sz': '4', 'color': '000000'}}

                self.set_cell_border(cell,border_style)
    def set_cell_border(self,cell, border_settings):
        """Устанавливает границы для одной ячейки."""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')

        for edge in border_settings:
            element = OxmlElement(f'w:{edge}')
            for key, value in border_settings[edge].items():
                element.set(qn(f'w:{key}'), str(value))
            tcBorders.append(element)

        tcPr.append(tcBorders)
    def apply_update_table_style(self, cells,table):
       
        table.autofit = False

        # Access the table XML element
        tbl = table._tbl
        tbl_pr = tbl.tblPr

        # Remove any existing tblW element
        existing_tblW = tbl_pr.xpath('./w:tblW')
        if existing_tblW:
            tbl_pr.remove(existing_tblW[0])

        # Create and append new tblW for 100% width (pct type, 5000 = 100%)
        tblW = OxmlElement('w:tblW')
        tblW.set(qn('w:type'), 'pct')
        tblW.set(qn('w:w'), '5000')
        tbl_pr.append(tblW)
        if cells:
            for cell in cells:
                for run in cell.paragraphs[0].runs:
                    run.font.name = self.config_service.get("documentFormatting.tables.family")  
                    run.font.size = Pt(self.config_service.get("documentFormatting.tables.size"))   
                    run.font.bold = self.config_service.get("documentFormatting.tables.bold")       
                    run.font.italic = self.config_service.get("documentFormatting.tables.style.italic")  
                    run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.tables.color")] 
                    paragraph_format  = cell.paragraphs[0].paragraph_format
    
                    paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.tables.paragraph_format.line_sp", 0)]
                    paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.indent_left", 0))      
                    paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.indent_right", 0))   
                    paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.tables.paragraph_format.first_line", 0))
                    paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.spacingBefore", 0))  # Space before paragraph
                    paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.tables.paragraph_format.spacingAfter", 0))
                    border_style = {
                        'top': {'val': 'single', 'sz': '4', 'color': '000000'},
                        'left': {'val': 'single', 'sz': '4', 'color': '000000'},
                        'bottom': {'val': 'single', 'sz': '4', 'color': '000000'},
                        'right': {'val': 'single', 'sz': '4', 'color': '000000'}}
    
                    self.set_cell_border(cell,border_style)


    def apply_formula_style(self, formula: Formula):
        
        table = formula.docx_element
        
                # --- Update formula in cell(0, 1) ---
        formula_cell = table.cell(0, 1)
       
        
        for run in  formula_cell.paragraphs[0].runs:
            # run.font.name = self.config_service.get("documentFormatting.formulas.family")  
            run.font.size = Pt(self.config_service.get("documentFormatting.formulas.size"))   
            run.font.bold = self.config_service.get("documentFormatting.formulas.style.bold")       
            run.font.italic = self.config_service.get("documentFormatting.formulas.style.italic")  
            run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.formulas.color")] 
        paragraph_format  = formula_cell.paragraphs[0].paragraph_format
        paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.formulas.lineSpacing", 0)]
        paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.indent_left", 0))      
        paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.indent_right", 0))   
        paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.formulas.paragraph_format.first_line", 0))
        paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.spacingBefore", 0))  # Space before paragraph
        paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.spacingAfter", 0))
        # --- Update caption in cell(0, 2) ---
        formula_cell.paragraphs[0].alignment = self.WD_PARAGRAPH_ALIGNMENT_MAP[self.config_service.get("documentFormatting.formulas.first.alignment", 0)]
        formula_cell.vertical_alignment =self.WD_CELL_VERTICAL_ALIGNMENT_MAP[self.config_service.get("documentFormatting.formulas.first.vertical", 0)]  
        
        caption_cell = table.cell(0, 2)
        for run in  caption_cell.paragraphs[0].runs:
            run.font.name = self.config_service.get("documentFormatting.formulas.family")  
            run.font.size = Pt(self.config_service.get("documentFormatting.formulas.size"))   
            run.font.bold = self.config_service.get("documentFormatting.formulas.style.bold")       
            run.font.italic = self.config_service.get("documentFormatting.formulas.style.italic")  
            run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.formulas.color")] 
        paragraph_format  = caption_cell.paragraphs[0].paragraph_format 
        caption_cell.paragraphs[0].alignment = self.WD_PARAGRAPH_ALIGNMENT_MAP[self.config_service.get("documentFormatting.formulas.second.alignment", 0)]
        caption_cell.vertical_alignment =self.WD_CELL_VERTICAL_ALIGNMENT_MAP[self.config_service.get("documentFormatting.formulas.second.vertical", 0)]  
        paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.formulas.lineSpacing", 0)]
        paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.indent_left", 0))      
        paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.indent_right", 0))   
        paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.formulas.paragraph_format.first_line", 0))
        paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.spacingBefore", 0))  # Space before paragraph
        paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.formulas.paragraph_format.spacingAfter", 0))
        section = self.doc.sections[0]
        page_width = section.page_width - section.left_margin - section.right_margin
        table.width = page_width

        table.autofit = False  # Фиксированные ширины
        table.allow_autofit = False
        table.columns[0].cells[0].width = int(page_width * 0.1)
        table.columns[1].cells[0].width = int(page_width * 0.8)
        table.columns[2].cells[0].width =int( page_width * 0.1)

            

    def apply_listing_style(self, listing: Listing):
        for run in listing.docx_caption_paragraph.runs:
            run.font.name = self.config_service.get("documentFormatting.listings.caption.family")  
            run.font.size = Pt(self.config_service.get("documentFormatting.listings.caption.size"))   
            run.font.bold = self.config_service.get("documentFormatting.listings.caption.style.bold")       
            run.font.italic = self.config_service.get("documentFormatting.listings.caption.style.italic")  
            run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.listings.caption.color")]

        paragraph_format  = listing.docx_caption_paragraph.paragraph_format
        paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.listings.caption.lineSpacing", 0)]
        paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.listings.paragraph_format.indent_left", 0))      
        paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.listings.paragraph_format.indent_right", 0))   
        paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.listings.paragraph_format.first_line", 0))
        paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.listings.caption.spacingBefore", 0))  # Space before paragraph
        paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.listings.caption.spacingAfter", 0))
        paragraph_format.alignment=self.WD_ALIGN_PARAGRAPH_MAP[self.config_service.get("documentFormatting.listings.caption.alignment", 0)]
        paragraph_format.keep_together=False
        paragraph_format.keep_with_next =False
        paragraph_format.widow_control = False
        for list_item in listing.docx_listing_paragraphs:
            for run in list_item.runs:
                run.font.name = self.config_service.get("documentFormatting.listings.family")  
                run.font.size = Pt(self.config_service.get("documentFormatting.listings.size"))   
                run.font.bold = self.config_service.get("documentFormatting.listings.style.bold")       
                run.font.italic = self.config_service.get("documentFormatting.listings.style.italic")  
                run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.listings.color")] 

            paragraph_format  = list_item.paragraph_format
            paragraph_format.keep_together=False
            paragraph_format.keep_with_next =False
            paragraph_format.widow_control = False
            paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.listings.lineSpacing", 0)]
            paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.listings.paragraph_format.indent_left", 0))      
            paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.listings.paragraph_format.indent_right", 0))   
            paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.listings.paragraph_format.first_line", 0))
            paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.listings.paragraph_format.spacingBefore", 0))  # Space before paragraph
            paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.listings.paragraph_format.spacingAfter", 0))

        p_bdr = OxmlElement('w:pBdr')  # Создаем элемент границ
            
        for paragraph in listing.docx_listing_paragraphs :
            p_pr = paragraph._element.get_or_add_pPr()
    
            suppress_elem = p_pr.find(qn('w:suppressAutoHyphens'))
            if suppress_elem is None:
                suppress_elem = parse_xml(r'<w:suppressAutoHyphens xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
                p_pr.append(suppress_elem)

            # Устанавливаем значение true для подавления переносов
            suppress_elem.set(qn('w:val'), 'true')
            p_bdr = OxmlElement('w:pBdr')
            for border_name in ['top', 'left', 'bottom', 'right']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'single')  # Тип линии
                border.set(qn('w:sz'), '4')         # Толщина (0.5 pt)
                border.set(qn('w:space'), '4')      # Отступ от текста
                border.set(qn('w:color'), '000000') # Цвет (#000000)
                p_bdr.append(border)

        p_pr.append(p_bdr)

    def apply_paragraph_style(self, paragraph: Paragraph):
        for run in paragraph.docx_element.runs:
                run.font.name = self.config_service.get("documentFormatting.paragraph.family")  
                run.font.size = Pt(self.config_service.get("documentFormatting.paragraph.size"))   
                run.font.bold = self.config_service.get("documentFormatting.paragraph.style.bold")       
                run.font.italic = self.config_service.get("documentFormatting.paragraph.style.italic")  
                run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.paragraph.color")] 

        paragraph_format  = paragraph.docx_element.paragraph_format
        paragraph_format.keep_together=False
        paragraph_format.keep_with_next =False
        paragraph_format.widow_control = False
        paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.paragraph.lineSpacing", 0)]
        if paragraph_format.line_spacing_rule is WD_LINE_SPACING.EXACTLY:
            paragraph_format.line_spacing=Pt(self.config_service.get("documentFormatting.general.font.line_space_count", 18))
        paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.indent_left", 0))      
        paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.indent_right", 0))   
        paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.first_line", 0))
        neighbor_before = self.neighbor_finder.get_before(paragraph)
        if isinstance(neighbor_before,Table): 
            paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
        else:
            paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBefore", 0))
            
        paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingAfter", 0))
        paragraph_format.alignment = self.WD_PARAGRAPH_ALIGNMENT_MAP[self.config_service.get("documentFormatting.paragraph.paragraph_format.alignment", 0)]
          

    def apply_after_table_if_table_exist(self,element):
        neighbor_before = self.neighbor_finder.get_before(element)
        if isinstance(element,Paragraph) and isinstance(neighbor_before,Table):
             element.docx_element.paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
        if isinstance(element,ListElement) and isinstance(neighbor_before,Table):
            element.docx_elements[0].paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
            
    def apply_after_table_if_table_exist_delete(self,element):
        neighbor_after = self.neighbor_finder.get_after(element)
        if isinstance(neighbor_after,Paragraph) and isinstance(neighbor_after,Table):
             neighbor_after.docx_element.paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
        if isinstance(neighbor_after,ListElement) and isinstance(neighbor_after,Table):
            neighbor_after.docx_elements[0].paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
            
        
    def apply_list_style(self, list_element: ListElement):
        neighbor_before = self.neighbor_finder.get_before(list_element)
        if isinstance(neighbor_before,Table): 
            list_element.docx_elements[0].paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBeforeTable", 0))
        else:
            list_element.docx_elements[0].paragraph_format.space_before = Pt(self.config_service.get("documentFormatting.paragraph.paragraph_format.spacingBefore", 0))
            
        for list_item in list_element.docx_elements:
            for run in list_item.runs:
                run.font.name = self.config_service.get("documentFormatting.lists.family")  
                run.font.size = Pt(self.config_service.get("documentFormatting.lists.size"))   
                run.font.bold = self.config_service.get("documentFormatting.lists.style.bold")       
                run.font.italic = self.config_service.get("documentFormatting.lists.style.italic")  
                run.font.color.rgb = self.COLOR_MAP[self.config_service.get("documentFormatting.lists.color")] 

            paragraph_format  = list_item.paragraph_format
            paragraph_format.keep_together=False
            paragraph_format.keep_with_next =False
            paragraph_format.widow_control = False
            
            paragraph_format.line_spacing_rule  =self.WD_LINE_SPACING_MAP[self.config_service.get("documentFormatting.lists.lineSpacing", 0)]
            if paragraph_format.line_spacing_rule is WD_LINE_SPACING.EXACTLY:
                paragraph_format.line_spacing=Pt(self.config_service.get("documentFormatting.general.font.line_space_count", 18))
            paragraph_format.left_indent = Pt(self.config_service.get("documentFormatting.lists.paragraph_format.indent_left", 0))      
            paragraph_format.right_indent = Pt(self.config_service.get("documentFormatting.lists.paragraph_format.indent_right", 0))   
            paragraph_format.first_line_indent = self.cm_to_pt(self.config_service.get("documentFormatting.lists.paragraph_format.first_line", 0))
            paragraph_format.space_after = Pt(self.config_service.get("documentFormatting.lists.paragraph_format.spacingAfter", 0))
            paragraph_format.alignment = self.WD_PARAGRAPH_ALIGNMENT_MAP[self.config_service.get("documentFormatting.paragraph.paragraph_format.alignment", 0)]
