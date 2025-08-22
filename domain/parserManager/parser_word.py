import re
from typing import List

import re

from docx import Document
from docx.oxml.ns import qn, nsmap
from lxml import etree

from elements.header import Header
from elements.formula import Formula
from elements.image import Image
from elements.list_element import ListElement
from elements.listing import Listing
from elements.table import Table, DataTable
from elements.paragraph import Paragraph

from domain.wordManager.tree_numbering import TreeNumbering
from elements.undefined_element import UndefinedElement
class ParserWord:
    def __init__(self,doc):
        self.doc=doc
    def is_header(self,paragraph):
        return paragraph.style.name.startswith('Heading')

    def get_header_level(self,paragraph):
        return int(paragraph.style.name.split()[1]) if paragraph.style.name.startswith('Heading') else 1

    def is_list_item(self,paragraph):
        # Check for styled lists first (original functionality)
        if paragraph.style is not None:
            style_name = paragraph.style.name.lower()
            if 'list' in style_name or 'bullet' in style_name or 'number' in style_name:
                return True

        # Check for dash-started lists with specific formatting
        if paragraph.text.strip().startswith('–') and paragraph.text.strip().endswith(';'):        
            return True
        if paragraph.text.strip().startswith('–') and paragraph.text.strip().endswith('.'):        
            return True

        return False

    def get_list_type(self,paragraph):
        if paragraph.style is None:
            return 'unordered'
        style_name = paragraph.style.name.lower()
        if 'bullet' in style_name:
            return 'unordered'
        elif 'number' in style_name:
            return 'ordered'
        return 'unordered'

    def is_table_caption(self,paragraph):
        text = paragraph.text.strip()
        pattern = r'^Таблица \d+\.\d+ – .+$'
        return bool(re.fullmatch(pattern, text))

    def is_image_caption(self,paragraph):
        text = paragraph.text.strip()
        pattern = r'^Рисунок \d+\.\d+ – .+$'
        return bool(re.fullmatch(pattern, text))

    def is_listing_caption(self,paragraph):
        text = paragraph.text.strip()
        pattern = r'^Листинг \d+\.\d+ – .+$'
        return bool(re.fullmatch(pattern, text))

    def is_formula_caption(self,paragraph):
        text = paragraph.text.strip()
        pattern = r'^\(\d+\.\d+\)$'
        return bool(re.fullmatch(pattern, text))

    def has_image(self,paragraph):
        has_pic = any(run.element.tag.endswith('pic') for run in paragraph.runs)
        has_inline = paragraph._element.findall('.//wp:inline', namespaces=nsmap)
        # print(f"Checking paragraph for image: has_pic={has_pic}, has_inline={len(has_inline) > 0}")
        return has_pic or has_inline

    def extract_image(self,paragraph):
        for run in paragraph.runs:
            if run.element.tag.endswith('pic'):
                blip = run.element.find('.//a:blip', namespaces=nsmap)
                if blip is not None:
                    rId = blip.get(qn('r:embed'))
                    if rId and rId in paragraph.part.related_parts:
                        image_part = paragraph.part.related_parts[rId]
                        return image_part.blob
        for inline in paragraph._element.findall('.//wp:inline', namespaces=nsmap):
            blip = inline.find('.//a:blip', namespaces=nsmap)
            if blip is not None:
                rId = blip.get(qn('r:embed'))
                if rId and rId in paragraph.part.related_parts:
                    image_part = paragraph.part.related_parts[rId]
                    return image_part.blob
        return b''

    def has_formula(self,paragraph):
        has_omath = paragraph._element.findall('.//m:oMath', namespaces=nsmap)
        has_omathpara = paragraph._element.findall('.//m:oMathPara', namespaces=nsmap)

        return has_omath or has_omathpara

    def extract_formula_and_caption(self,paragraph):
        formula_xml = []
        omath = paragraph._element.findall('.//m:oMath', namespaces=nsmap)
        omathpara = paragraph._element.findall('.//m:oMathPara', namespaces=nsmap)
        for elem in omath + omathpara:
            formula_xml.append(etree.tostring(elem, encoding='unicode'))
        formula = ''.join(formula_xml)
        caption = None
        text = paragraph.text.strip()
        caption_match = re.search(r'\(\d+\.\d+\)', text)
        if caption_match:
            caption = caption_match.group(0)
        return formula, caption

    def is_listing(self,paragraph):
        # Проверяем наличие границ (w:pBdr) в свойствах параграфа
        pPr = paragraph._element.find('.//w:pPr', namespaces=nsmap)
        has_borders = pPr is not None and pPr.find('.//w:pBdr', namespaces=nsmap) is not None
        is_listing_style = paragraph.style.name == 'Listing'
        return has_borders or is_listing_style

    def parse_document(self):
        tree = []
        stack = [tree]
        current_level = 0
        pending_table_caption = None
        paragraph_index = 0
        table_index = 0
        elements = list(self.doc.element.body)

        i = 0
        while i < len(elements):
            element = elements[i]
            if element.tag.endswith('p'):
                if paragraph_index >= len(self.doc.paragraphs):
                    parent = stack[-1]
                    parent.append(UndefinedElement(str(element), element))
                    i += 1
                    continue
                paragraph = self.doc.paragraphs[paragraph_index]
                paragraph_index += 1
                if self.is_header(paragraph):
                    level = self.get_header_level(paragraph)
                    text = paragraph.text.strip()
                    if not text:
                        i += 1
                        continue
                    while len(stack) > 1 and current_level >= level:
                        stack.pop()
                        current_level -= 1
                    parent = stack[-1]
                    new_header = Header(level, text, [], paragraph, "")
                    parent.append(new_header)
                    stack.append(new_header.children)
                    current_level = level
                    i += 1
                elif self.is_list_item(paragraph):
                    if paragraph.style is not None:  # Original styled list handling
                        list_type = self.get_list_type(paragraph)
                        list_items = [paragraph.text.strip() or "(empty)"]
                        list_paragraphs = [paragraph]
                        j = i + 1
                        next_paragraph_index = paragraph_index
                        while (j < len(elements) and elements[j].tag.endswith('p') and 
                               next_paragraph_index < len(self.doc.paragraphs) and
                               self.is_list_item(self.doc.paragraphs[next_paragraph_index]) and 
                               self.get_list_type(self.doc.paragraphs[next_paragraph_index]) == list_type):
                            next_text = self.doc.paragraphs[next_paragraph_index].text.strip()
                            list_items.append(next_text or "(empty)")
                            list_paragraphs.append(self.doc.paragraphs[next_paragraph_index])
                            next_paragraph_index += 1
                            j += 1
                        parent = stack[-1]
                        parent.append(ListElement(list_items, list_type, list_paragraphs))
                        i = j
                        paragraph_index = next_paragraph_index
                    else:  # Handle dash-started lists
                        list_items = []
                        list_paragraphs = []
                        current_index = i
                        current_para_index = paragraph_index
                        list_complete = False

                        # Collect all consecutive dash-started paragraphs until we find the terminating one
                        while (current_index < len(elements) and elements[current_index].tag.endswith('p') and
                               current_para_index < len(self.doc.paragraphs)):
                            current_para = self.doc.paragraphs[current_para_index]
                            text = current_para.text.strip()

                            if not text.startswith('–'):
                                break

                            list_items.append(text or "(empty)")
                            list_paragraphs.append(current_para)

                            # Check if this is the terminating item (ends with .)
                            if text.endswith('.'):
                                list_complete = True
                                break

                            current_index += 1
                            current_para_index += 1

                        # Only create list if we found proper termination
                        if list_complete:
                            parent = stack[-1]
                            parent.append(ListElement(list_items, "dash", list_paragraphs))
                            i = current_index + 1
                            paragraph_index = current_para_index + 1
                        else:
                            # If not a complete list, treat as regular paragraphs
                            parent = stack[-1]
                            parent.append(Paragraph(paragraph.text, paragraph))
                            i += 1
                            paragraph_index += 1

                elif self.is_table_caption(paragraph):
                    pending_table_caption = (paragraph.text.strip(), paragraph)
                    i += 1
                elif self.has_image(paragraph):
                    image_data = self.extract_image(paragraph)
                    caption = None
                    caption_paragraph = None
                    if (i + 1 < len(elements) and elements[i + 1].tag.endswith('p') and 
                        paragraph_index < len(self.doc.paragraphs) and
                        self.is_image_caption(self.doc.paragraphs[paragraph_index])):
                        caption = self.doc.paragraphs[paragraph_index].text.strip()
                        caption_paragraph = self.doc.paragraphs[paragraph_index]
                        paragraph_index += 1
                        i += 1
                    parent = stack[-1]
                    parent.append(Image(image_data, caption, paragraph, caption_paragraph, ""))
                    i += 1
                elif self.is_listing(paragraph):
                    listing_texts = [paragraph.text.strip() or "(empty)"]
                    listing_paragraphs = [paragraph]
                    j = i + 1
                    next_paragraph_index = paragraph_index
                    while (j < len(elements) and elements[j].tag.endswith('p') and 
                           next_paragraph_index < len(self.doc.paragraphs) and
                           self.is_listing(self.doc.paragraphs[next_paragraph_index]) and
                           not self.is_listing_caption(self.doc.paragraphs[next_paragraph_index])):
                        next_text = self.doc.paragraphs[next_paragraph_index].text.strip()
                        listing_texts.append(next_text or "(empty)")
                        listing_paragraphs.append(self.doc.paragraphs[next_paragraph_index])
                        next_paragraph_index += 1
                        j += 1
                    text = '\n'.join(listing_texts)
                    caption = None
                    caption_paragraph = None
                    if (j < len(elements) and elements[j].tag.endswith('p') and 
                        next_paragraph_index < len(self.doc.paragraphs) and
                        self.is_listing_caption(self.doc.paragraphs[next_paragraph_index])):
                        caption = self.doc.paragraphs[next_paragraph_index].text.strip()
                        caption_paragraph = self.doc.paragraphs[next_paragraph_index]
                        next_paragraph_index += 1
                        j += 1
                    parent = stack[-1]
                    parent.append(Listing(text, caption, listing_paragraphs, caption_paragraph, ""))
                    i = j
                    paragraph_index = next_paragraph_index
                else:
                    text = paragraph.text.strip()
                    if text:
                        parent = stack[-1]
                        parent.append(Paragraph(text, paragraph))
                    i += 1
            elif element.tag.endswith('tbl'):
                if table_index < len(self.doc.tables):
                    table = self.doc.tables[table_index]
                    table_index += 1
                    is_formula_table = (
                        len(table.rows) == 1 and
                        len(table.columns) == 3 and
                        self.has_formula(table.cell(0, 1).paragraphs[0])
                    )
                    if is_formula_table:
                        # Извлекаем формулу и подпись
                        formula, _ = self.extract_formula_and_caption(table.cell(0, 1).paragraphs[0])
                        caption = table.cell(0, 2).text.strip() or None

                        parent = stack[-1]
                        parent.append(Formula(formula, caption, table,
                                              table.cell(0, 2).paragraphs[0], caption))
                    else: 
                        caption, caption_paragraph = pending_table_caption if pending_table_caption else (None, None)
                        pending_table_caption = None

                        # Парсим таблицу с учётом объединений
                        data = self.parse_table_with_merging(table)

                        parent = stack[-1]
                        parent.append(Table(data, caption, table, caption_paragraph, ""))
                else:
                    parent = stack[-1]
                    parent.append(UndefinedElement(str(element), element))
                i += 1
            else:
                parent = stack[-1]
                parent.append(UndefinedElement(str(element), element))
                i += 1

        return tree
    


    def parse_table_with_merging(self, table) -> List[List[DataTable]]:
        """Парсит таблицу, учитывая объединённые и разделённые ячейки"""
        data = []
        ns = {"w14": "http://schemas.microsoft.com/office/word/2010/wordml"}
      
        num_rows = len(table.rows)
        num_cols = len(table.columns)

        visited = [[(-1,-1) for _ in range(num_cols)] for _ in range(num_rows)]

        for i, row in enumerate(table.rows):
            data_row = []
            for j, cell in enumerate(row.cells):
                # Получаем свойства ячейки из XML
                if visited[i][j] !=(-1,-1):
                    cell_text = '\n'.join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                    x, y = visited[i][j]
                    cell_data = DataTable(
                    data=cell_text,
                    horizontal=cell.paragraphs[0].alignment if cell.paragraphs else None,
                    vertical=cell.vertical_alignment,
                    parent_row=x,
                    parent_col=y)
                    data_row.append(cell_data)
                    continue
                tc_pr = cell._element.tcPr
                grid_span = None

                if tc_pr is not None:
                    grid_span = next((e for e in tc_pr if e.tag.endswith('gridSpan')), None)

                colspan = int(grid_span.get(qn("w:val"))) if grid_span is not None else 1

                paraId=None
                if cell.paragraphs:
                    paraId = cell.paragraphs[0]._element.get(f"{{{ns['w14']}}}paraId")
                # Определяем rowspan (если ячейка объединена вертикально)
                    rowspan = 1
                    for k in range(i + 1, len(table.rows)):
                        try:
                            next_cell = table.cell(k, j)
                            next_paraId=None
                            if next_cell.paragraphs:
                                next_paraId = next_cell.paragraphs[0]._element.get(f"{{{ns['w14']}}}paraId")

                     
                            if  next_paraId==paraId:
                                if len(table.rows)==k+1:
                                    bottom_left=(k,j)
                                    bottom_right = (k,j+colspan-1)
                                    top_left = (i,j)
                                    top_right = (i,j+colspan-1)
                                    visited[top_right[0]][top_right[1]] = (i,j)
                                    visited[top_left[0]][top_left[1]] = (i,j)
                                    visited[bottom_right[0]][bottom_right[1]] = (i,j)
                                    visited[bottom_left[0]][bottom_left[1]] = (i,j)
                                    for x in range(top_left[0], bottom_left[0] + 1):
                                        for y in range(top_left[1], top_right[1] + 1):
                                            visited[x][y] = (i,j)

                            else:
                                
                                bottom_left=(k-1,j)
                                bottom_right = (k-1,j+colspan-1)
                                top_left = (i,j)
                                top_right = (i,j+colspan-1)
                                visited[top_right[0]][top_right[1]] = (i,j)
                                visited[top_left[0]][top_left[1]] = (i,j)
                                visited[bottom_right[0]][bottom_right[1]] = (i,j)
                                visited[bottom_left[0]][bottom_left[1]] = (i,j)
                                for x in range(top_left[0], bottom_left[0] + 1):
                                    for y in range(top_left[1], top_right[1] + 1):
                                        visited[x][y] = (i,j)
                                break
                        except IndexError:
                            break

                # Получаем текст ячейки
                cell_text = '\n'.join(p.text.strip() for p in cell.paragraphs if p.text.strip())

                # Создаём DataTable
                cell_data = DataTable(
                    data=cell_text,
                    horizontal=cell.paragraphs[0].alignment if cell.paragraphs else None,
                    vertical=cell.vertical_alignment,
                    colspan=colspan,
                    parent_row=i,
                    parent_col=j
                )
                data_row.append(cell_data)

            data.append(data_row)

        return data
    def print_tree(self,elements, indent=0, prefix=""):
        for i, elem in enumerate(elements):
            is_last = i == len(elements) - 1
            branch = "└── " if is_last else "├── "
            indent_str = prefix + branch
            content = ""

            if isinstance(elem, Header):
                content = f"Header (Level {elem.level}, {elem.text_number_of_element}): {elem.text[:50]}{'...' if len(elem.text) > 50 else ''}"
            elif isinstance(elem, Paragraph):
                content = f"Paragraph: {elem.text[:50]}{'...' if len(elem.text) > 50 else ''}"
            elif isinstance(elem, ListElement):
                first_item = elem.items[0][:47] + '...' if len(elem.items) > 0 and len(elem.items[0]) > 47 else elem.items[0] if len(elem.items) > 0 else 'No items'
                content = f"List ({elem.list_type}, {len(elem.items)} items): {first_item}"
            elif isinstance(elem, Image):
                image_size = f"{len(elem.image_data)} bytes" if elem.image_data else "No data"
                content = f"Image ({elem.text_number_of_element}): {elem.caption[:44] if elem.caption else 'No caption'} ({image_size})"
                content = content[:50] + '...' if len(content) > 50 else content
            elif isinstance(elem, Table):
                first_cell = elem.data[0][0].data[:41] if elem.data and elem.data[0] else 'No data'
                content = f"Table ({elem.text_number_of_element}, {len(elem.data)} rows): {elem.caption[:35] if elem.caption else 'No caption'}, {first_cell}"
                content = content[:50] + '...' if len(content) > 50 else content
            elif isinstance(elem, Formula):
                formula_preview = elem.formula[:41] if elem.formula else 'No formula'
                content = f"Formula ({elem.text_number_of_element}): {elem.caption[:35] if elem.caption else 'No caption'}, {formula_preview}"
                content = content[:50] + '...' if len(content) > 50 else content
            elif isinstance(elem, Listing):
                listing_preview = elem.text[:41] if elem.text else 'No text'
                content = f"Listing ({elem.text_number_of_element}): {elem.caption[:35] if elem.caption else 'No caption'}, {listing_preview}"
                content = content[:50] + '...' if len(content) > 50 else content
            elif isinstance(elem, UndefinedElement):
                content = f"Undefined: {elem.content[:50]}{'...' if len(elem.content) > 50 else ''}"

            print(f"{indent_str}{content}")

            if isinstance(elem, Header) and elem.children:
                new_prefix = prefix + ("    " if is_last else "│   ")
                self.print_tree(elem.children, indent + 1, new_prefix)



# Функция для красивого вывода дерева

# Пример использования
if __name__ == "__main__":
    
    doc_path = "../../tst.docx"
    doc = Document(doc_path)
    parser_wr= ParserWord(doc)
    elements_tree = parser_wr.parse_document()
    TreeNumbering(elements_tree).assign_numbers()
    print("Document Structure:")
    parser_wr.print_tree(elements_tree)
  