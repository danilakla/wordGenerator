from dataclasses import dataclass, field
from typing import List, Optional
from docx.text.paragraph import Paragraph as DocxParagraph
from docx.table import Table as DocxTable
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

class Element:
    pass

# Заголовок
@dataclass
class Header(Element):
    level: int
    text: str
    children: List['Element'] = field(default_factory=list)
    docx_element: Optional[DocxParagraph] = None
    text_number_of_element: str = ""

# Параграф
@dataclass
class Paragraph(Element):
    text: str
    docx_element: Optional[DocxParagraph] = None

# Список
@dataclass
class ListElement(Element):
    items: List[str]
    list_type: str
    docx_elements: List[DocxParagraph] = field(default_factory=list)
    @property
    def list_type(self) -> str:
        return self._list_type

    @list_type.setter
    def list_type(self, value: str):
        self._list_type = value
        # Concatenate the list_type with each item in items
        self.items = [f"{value} {item}" for item in self.items]

# Рисунок
@dataclass
class Image(Element):
    image_data: bytes
    caption: str
    docx_image_paragraph: Optional[DocxParagraph] = None
    docx_caption_paragraph: Optional[DocxParagraph] = None
    text_number_of_element: str = ""

@dataclass
class DataTable:
    data: str
    horizontal: WD_PARAGRAPH_ALIGNMENT
    vertical: WD_CELL_VERTICAL_ALIGNMENT
    colspan: int =0
    parent_col:Optional[int] =None
    parent_row:Optional[int] =None

@dataclass
class Table:
    data: List[List[DataTable]] = field(default_factory=list)
    caption: str = ""
    docx_table: Optional[DocxTable] = None
    docx_caption_paragraph: Optional[DocxParagraph] = None
    text_number_of_element: str = ""
    
    def add_column_right(self):
        """
        Добавляет колонку справа в существующую матрицу (in-place).
        """
        if not self.data:
            self.data.append([DataTable()])
            return

        num_cols = len(self.data[0])

        for row_idx, row in enumerate(self.data):
            # Добавляем новую ячейку в конец строки
            new_cell = DataTable(parent_col=num_cols, parent_row=row_idx,data='',horizontal=None, vertical=None)
            row.append(new_cell)
    def add_column_left(self):
        """
        Добавляет колонку слева в существующую матрицу (in-place).
        """
        if not self.data:
            self.data.append([DataTable()])
            return

        for row_idx, row in enumerate(self.data):
            # Вставляем новую ячейку в начало строки
            new_cell = DataTable(parent_col=0, parent_row=row_idx,data='',horizontal=None, vertical=None)
            row.insert(0, new_cell)

            # Обновляем parent_col для всех остальных ячеек в строке
            for col_idx in range(1, len(row)):
                if row[col_idx].parent_col >= 0:
                    row[col_idx].parent_col += 1
    def insert_row_above(self, row_idx: int):
        """
        Добавляет новую строку над строкой row_idx (новая строка встанет по индексу row_idx,
        исходная "целевая" строка сдвинется вниз на row_idx+1).
        Проверяем строку над текущей (если есть), чтобы понять, совпадают ли parent_col/parent_row
        между верхней и текущей ячейками. Если совпадают — наследуем parent. Иначе создаём новое
        множество parent (первая ячейка группы становится родителем) с parent_row = row_idx (индекс новой строки).
        """
        old_row = self.data[row_idx]
        new_row: List[DataTable] = []

        prev_row = self.data[row_idx - 1] if row_idx - 1 >= 0 else None
        self.data.insert(row_idx, new_row)
        col_pointer = 0
        while col_pointer < len(old_row):
            cell = old_row[col_pointer]
            colspan = cell.colspan or 1

            top_cell = None
            if prev_row and col_pointer < len(prev_row):
                top_cell = prev_row[col_pointer]

            if top_cell and top_cell.parent_col == cell.parent_col and top_cell.parent_row == cell.parent_row:
                actial_col=len(new_row)
                iter_col=actial_col
                for _ in range(colspan):
                    new_row.append(DataTable(
                        data="",
                        horizontal=cell.horizontal,
                        vertical=cell.vertical,
                        colspan=colspan,
                        parent_col=cell.parent_col,
                        parent_row=cell.parent_row
                    ))
                    actial_row=row_idx+1
                    temp_col=cell.parent_col
                    temp_row=cell.parent_row
                    exsitRows=True
                    while cell.parent_col== temp_col and cell.parent_row==temp_row:
                        
                        
                        actial_row+=1
                        if actial_row>=len(self.data):
                            exsitRows=False
                            break    
                        temp_col=self.data[actial_row][iter_col].parent_col 
                        temp_row=self.data[actial_row][iter_col].parent_row
                    if exsitRows:
                        for l in range(actial_row,len(self.data)):
                            self.data[l][iter_col].parent_row+=1
                    iter_col+=1
            else:
                parent_col = len(new_row)
                parent_row = row_idx  # индекс НОВОЙ строки (так как вставляем сверху)
                iter_col=parent_col
                for j in range(colspan):
                    new_row.append(DataTable(
                        data="",
                        horizontal=cell.horizontal,
                        vertical=cell.vertical,
                        colspan=colspan,
                        parent_col=parent_col,
                        parent_row=parent_row
                    ))
                    for l in range(parent_row+1,len(self.data)):
                        self.data[l][iter_col].parent_row+=1
                    iter_col+=1

            col_pointer += colspan
        self.data.pop(row_idx)
        self.data.insert(row_idx, new_row)
    def insert_row_under(self, row_idx: int):
        """
        Добавляет новую строку под строкой row_idx.
        Проверяем строку под текущей (если есть), чтобы понять, совпадают ли parent_col/parent_row.
        Если совпадают, то новая ячейка наследует эти parent.
        Иначе создаём новое множество parent (первая ячейка группы становится родителем).
        """
        old_row = self.data[row_idx]
        new_row: List[DataTable] = []

        next_row = self.data[row_idx + 1] if row_idx + 1 < len(self.data) else None
        self.data.insert(row_idx + 1, [])

        col_pointer = 0
        while col_pointer < len(old_row):
            cell = old_row[col_pointer]
            colspan = cell.colspan or 1

            # Проверяем ячейку снизу (если есть)
            bottom_cell = None
            if next_row and col_pointer < len(next_row):
                bottom_cell = next_row[col_pointer]

            if bottom_cell and bottom_cell.parent_col == cell.parent_col and bottom_cell.parent_row == cell.parent_row:
                # наследуем parent из связки сверху/снизу
                
                actial_col=len(new_row)
                iter_col=actial_col
                for _ in range(colspan):
                    new_row.append(DataTable(
                        data="",
                        horizontal=cell.horizontal,
                        vertical=cell.vertical,
                        colspan=colspan,
                        parent_col=cell.parent_col,
                        parent_row=cell.parent_row
                    ))
                    actial_row=row_idx+2
                    temp_col=cell.parent_col
                    temp_row=cell.parent_row
                    exsitRows=True
                    while cell.parent_col== temp_col and cell.parent_row==temp_row:
                        
                        
                        actial_row+=1
                        if actial_row>=len(self.data):
                            exsitRows=False
                            break    
                        temp_col=self.data[actial_row][iter_col].parent_col 
                        temp_row=self.data[actial_row][iter_col].parent_row
                    if exsitRows:
                        for l in range(actial_row,len(self.data)):
                            self.data[l][iter_col].parent_row+=1
                    iter_col+=1
                
            else:
                # создаём новое множество parent
                parent_col = len(new_row)
                parent_row = row_idx + 1  # индекс новой строки
                iter_col=parent_col
                for j in range(colspan):
                    new_row.append(DataTable(
                        data="",
                        horizontal=cell.horizontal,
                        vertical=cell.vertical,
                        colspan=colspan,
                        parent_col=parent_col,
                        parent_row=parent_row
                    ))
                    for l in range(parent_row+1,len(self.data)):
                        self.data[l][iter_col].parent_row+=1
                    iter_col+=1

            col_pointer += colspan

        # Вставляем строку под исходной
        self.data.pop(row_idx+1)
        self.data.insert(row_idx + 1, new_row)

        
    def update_docx_table(self):
        """Обновляет docx-таблицу на основе данных в `self.data`"""
        for i, row in enumerate(self.data):
            for j, cell_data in enumerate(row):
                try:
                    cell = self.docx_table.cell(i, j)
                except IndexError:
                    raise IndexError(f"Cell ({i}, {j}) is outside the table bounds")

                # Очистка существующих параграфов
                for paragraph in cell.paragraphs[1:]:
                    paragraph.clear()
                
                if not cell.paragraphs:
                    cell.add_paragraph(cell_data.data)
                else:
                    if cell.paragraphs[0].runs:
                        for run in cell.paragraphs[0].runs:
                            run.text = ""
                        cell.paragraphs[0].runs[0].text = cell_data.data
                    else:
                        cell.paragraphs[0].text = cell_data.data

                # Обновление выравнивания
                cell.paragraphs[0].alignment = cell_data.horizontal
                cell.vertical_alignment = cell_data.vertical

               
    


# Формула
@dataclass
class Formula(Element):
    formula: str
    caption: str
    docx_element: Optional[DocxTable] = None
    docx_caption_paragraph: Optional[DocxParagraph] = None
    text_number_of_element: str = ""

# Листинг
@dataclass
class Listing(Element):
    text: str
    caption: str
    docx_listing_paragraphs: List[DocxParagraph] = field(default_factory=list)
    docx_caption_paragraph: Optional[DocxParagraph] = None
    text_number_of_element: str = ""

# Неизвестный элемент
@dataclass
class UndefinedElement(Element):
    content: str
    docx_element: Optional[object] = None