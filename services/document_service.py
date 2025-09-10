# document_service.py
from datetime import datetime
from typing import Dict, Optional, List
from docx import Document
from io import BytesIO
from PIL import Image as PILImage
import copy
import os

from elements.header import Header
from elements.paragraph import Paragraph
from elements.table import Table, DataTable
from elements.image import Image
from elements.listing import Listing
from elements.list_element import ListElement
from elements.formula import Formula
from elements.element import Element
from domain.wordManager.word_document_manager import WordDocumentManager

class DocumentService:
    def __init__(self):
        self.university=""
        self.manager = WordDocumentManager()
        self.doc_path = None
        self.save_folder = None

    def set_save_folder(self, folder: str):
        self.save_folder = folder

    def create_new_document(self, name: str):
        self.doc_path = os.path.join(self.save_folder, f"{name}.docx")
        doc = Document()
        doc.save(self.doc_path)
        self.manager.load_document(self.doc_path,self.university)

    def load_document(self, path: str):
        directory, filename = os.path.split(path)
        date_str = datetime.now().strftime("%Y.%m.%d")
        new_filename =  date_str +"_новый_документ"+"_"+filename
        new_path = os.path.join(directory, new_filename)

        self.doc_path = new_path
        self.manager.load_document(path,self.university)

    def save_document(self):
        self.manager.save_document(self.doc_path)
        
    def temp_save_d(self,temp_d):
        self.manager.save_document(temp_d)
        
        
    def emergency_save(self):
        return self.manager.emergency_save()
    

    def get_tree(self) -> List[Element]:
        return self.manager.tree

    def get_parent(self, target: Element) -> Optional[Element]:
        def find_parent(elements: List[Element], current_parent: Optional[Element] = None) -> Optional[Element]:
            for elem in elements:
                if elem == target:
                    return current_parent
                if hasattr(elem, 'children'):
                    parent = find_parent(elem.children, elem)
                    if parent is not None:
                        return parent
            return None
        return find_parent(self.get_tree())

    def add_header(self, parent: Optional[Element], text: str, level: int, after: Optional[Element] = None):
        header = Header(level=level, text=text)
        self.manager.add_element(parent, header, after)

    def add_paragraph(self, parent: Optional[Element], text: str, after: Optional[Element] = None):
        paragraph = Paragraph(text=text)
        self.manager.add_element(parent, paragraph, after)
        return paragraph

    def add_table(self, parent: Optional[Element], data: List[List[DataTable]], caption: str, after: Optional[Element] = None):
        table = Table(data=data, caption=caption)
        self.manager.add_element(parent, table, after)

    def add_image(self, parent: Optional[Element], image_path: str, caption: str, after: Optional[Element] = None):
        with open(image_path, 'rb') as f:
            image_data = f.read()
        image = Image(image_data=image_data, caption=caption)
        self.manager.add_element(parent, image, after)

    def add_listing(self, parent: Optional[Element], text: str, caption: str, after: Optional[Element] = None):
        listing = Listing(text=text, caption=caption)
        self.manager.add_element(parent, listing, after)

    def add_list(self, parent: Optional[Element], items: List[str], list_type: str, after: Optional[Element] = None):
        list_elem = ListElement(items=items, list_type=list_type)
        self.manager.add_element(parent, list_elem, after)

    def add_formula(self, parent: Optional[Element], formula: str, caption: str, after: Optional[Element] = None):
        formula_elem = Formula(formula=formula, caption=caption)
        self.manager.add_element(parent, formula_elem, after)

    def remove_element(self, element: Element):
        self.manager.remove_element(element)

    def update_element(self, element: Element, new_data: Dict):
        if isinstance(element, Header):
            new_header = Header(level=element.level, text=new_data['text'])
            self.manager.update_element(element, new_header)
        elif isinstance(element, Paragraph):
            new_para = Paragraph(text=new_data['text'])
            self.manager.update_element(element, new_para)
        elif isinstance(element, Table):
            new_table = Table(data=new_data['data'], caption=new_data['caption'])
            self.manager.update_element(element, new_table)
        elif isinstance(element, Image):
            caption = new_data.get('caption', element.caption)
            image_data = element.image_data
            if 'image_path' in new_data and new_data['image_path']:
                with open(new_data['image_path'], 'rb') as f:
                    image_data = f.read()
            new_image = Image(image_data=image_data, caption=caption)
            self.manager.update_element(element, new_image)
        elif isinstance(element, Listing):
            new_listing = Listing(text=new_data['text'], caption=new_data['caption'])
            self.manager.update_element(element, new_listing)
        elif isinstance(element, ListElement):
            new_list = ListElement(items=new_data['items'], list_type=new_data['list_type'])
            self.manager.update_element(element, new_list)
        elif isinstance(element, Formula):
            formula = new_data['formula']
            caption = new_data.get('caption', element.caption)
            new_formula = Formula(formula=formula, caption=caption)
            self.manager.update_element(element, new_formula)

    def move_element(self, element: Element, new_parent: Optional[Element], after: Optional[Element] = None):
        self.manager.move_element_word(new_parent, element, after)

    def get_doc_path(self) -> str:
        return self.doc_path