# Создание таблицы
from docx import Document

from elements import DataTable, Table

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
# Создание таблицы
doc = Document()
docx_table = doc.add_table(rows=3, cols=3)
table = Table(
    data=[],
    caption="Пример",
    docx_table=docx_table
)

# Заполнение данными (автоматически синхронизируется)
for i in range(3):
    for j in range(3):
        cell_data = table.data[i][j]
        cell_data.data = f"{i}-{j}"
        cell_data.horizontal = WD_PARAGRAPH_ALIGNMENT.CENTER
        cell_data.vertical = WD_CELL_VERTICAL_ALIGNMENT.CENTER

# Объединение ячеек
table.merge_cells(0, 0, 1, 1)  # Объединяем 4 ячейки

# Разделение ячейки
table.split_cell(2, 2, 2, 2)   # Разделяем на 4 ячейки

# Просмотр структуры
table.print_debug_info()

# Сохранение изменений
# table.update_docx_table()
doc.save("table_example.docx")